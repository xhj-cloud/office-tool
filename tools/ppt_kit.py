"""
PPT 设计工具包（赛博深蓝 · 暗色极客风）
"一次做好" 质量闭环：python-pptx 精确像素绘制（文字可编辑）
+ Pillow 同坐标预览（read_image 逐页目检）+ 溢出警告。

2026-10-06 自 ~/Desktop/PPT工坊 集成，已验证可逐字节复现 DeepSeek 交付版 PPTX。
"""

import json
import math
import os
import random
import tempfile
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

from .json_repair import safe_parse_json

# ── 画布与单位 ─────────────────────────────────────────────
W, H = 1600, 900          # px 画布（13.333in × 7.5in, 16:9）
EMU_PER_PX = 7620
SLIDE_W_EMU, SLIDE_H_EMU = 12192000, 6858000
NO_GRID_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
GHOSTS = ["candles", "target", "bars", "globe", "gear"]

# ── 字体（macOS 系统字体；缺失时回退默认并给出警告） ────────
FONT_CJK = "PingFang SC"
FONT_MONO = "Menlo"
CJK_TTC = "/System/Library/Fonts/Hiragino Sans GB.ttc"   # idx 0 常规 / 2 粗
MONO_TTC = "/System/Library/Fonts/Menlo.ttc"             # idx 0 常规 / 1 粗
_FONT_WARN = []
if not os.path.exists(CJK_TTC) or not os.path.exists(MONO_TTC):
    _FONT_WARN.append("预览字体缺失（非 macOS?）：Hiragino Sans GB / Menlo，预览仅示意，PPTX 本身不受影响")


def _E(px):
    return int(round(px * EMU_PER_PX))


# ── 风格色板 ────────────────────────────────────────────────
STYLES = {
    "cyber": {  # 赛博深蓝
        "bg": "0B1220", "tag": "7FA8C9", "title": "5EA0FF", "body": "D6E2F0",
        "cyan": "22D3EE", "dim": "8CA3BE",
        "chipf": "0F1A2E", "chipb": "1E5B73", "chiphl": "12233C",
        "cell": "111B2E", "th": "1E2A44", "tr0": "101A2C", "tr1": "0D1626",
        "chipname": "BFD9E8",
        "bg_rgb": (11, 18, 32), "cyan_rgb": (34, 211, 238), "blue_rgb": (94, 160, 255),
    },
}


def _mix(bg, c, a):
    if a >= 255:
        return c
    return tuple(int(bg[i] + (c[i] - bg[i]) * a / 255) for i in range(3))


# ═══════════════════════════════════════════════════════════
# 背景装饰层
# ═══════════════════════════════════════════════════════════

def _mono_font(size, bold=False):
    try:
        return ImageFont.truetype(MONO_TTC, size, index=1 if bold else 0)
    except Exception:
        return ImageFont.load_default()


def _icon_test(kind, x, y):
    dx, dy = x - 128, y - 128
    r = math.hypot(dx, dy)
    if kind == "candles":
        for (wx, wt, wb, bt, bb) in [(55, 70, 180, 95, 165), (95, 60, 150, 75, 135),
                                     (135, 80, 190, 100, 180), (175, 50, 140, 60, 120),
                                     (215, 40, 120, 45, 105)]:
            if abs(x - wx) <= 1 and wt <= y <= wb:
                return True
            if abs(x - wx) <= 8 and bt <= y <= bb:
                return True
        if 30 <= x <= 226 and 222 <= y <= 226:
            return True
        return False
    if kind == "target":
        if 98 <= r <= 102 or 63 <= r <= 67 or 28 <= r <= 32 or r <= 10:
            return True
        for (x0, y0, x1, y1) in [(126, 14, 130, 46), (126, 210, 130, 242),
                                 (14, 126, 46, 130), (210, 126, 242, 130)]:
            if x0 <= x <= x1 and y0 <= y <= y1:
                return True
        return False
    if kind == "bars":
        for bx, bh in [(45, 65), (90, 105), (135, 155), (180, 180)]:
            if bx - 14 <= x <= bx + 14 and 215 - bh <= y <= 215:
                return True
        if 30 <= x <= 226 and 211 <= y <= 215:
            return True
        return False
    if kind == "globe":
        if 93 <= r <= 97:
            return True
        if abs(dy) <= 2.5 and (dx / 95) ** 2 + (dy / 95) ** 2 <= 1:
            return True
        v = (dx / 45) ** 2 + (dy / 95) ** 2
        return 0.75 <= v <= 1.08
    if kind == "gear":
        if r <= 26:
            return False
        if r <= 62:
            BOLT = [(134, 92), (112, 132), (126, 132), (120, 164), (146, 120), (130, 120)]
            s = False
            j = 5
            for i in range(6):
                xi, yi = BOLT[i]
                xj, yj = BOLT[j]
                if ((yi > y) != (yj > y)) and (x < (xj - xi) * (y - yi) / (yj - yi) + xi):
                    s = not s
                j = i
            return not s
        if r <= 98:
            ang = math.degrees(math.atan2(dy, dx)) % 45
            return ang < 18 or ang > 27
        return False
    raise ValueError(f"未知幽灵图标: {kind}")


def _ghost_icon(img, kind, cx, cy, f=1.3, a=16):
    d = ImageDraw.Draw(img)
    r = 115 * f
    for y in range(int(cy - r), int(cy + r) + 1):
        for x in range(int(cx - r), int(cx + r) + 1):
            xd, yd = (x - cx) / f + 128, (y - cy) / f + 128
            if 0 <= xd <= 256 and 0 <= yd <= 256 and _icon_test(kind, xd, yd):
                d.point((x, y), fill=_mix(STYLES["cyber"]["bg_rgb"], STYLES["cyber"]["cyan_rgb"], a))


def _chip(img, cx, cy, f=200 / 512):
    def t(xd, yd):
        dx, dy = abs(xd - 256), abs(yd - 256)
        m = max(dx, dy)
        if 128 <= m <= 156 or 48 <= m <= 78 or m <= 24:
            return True
        pos = [116 + 280 * (i + 0.5) / 6 for i in range(6)]
        pw, L = 14, 48
        for p in pos:
            if abs(yd - p) <= pw and ((dx < 116 and dx > 116 - L) or (dx > 396 and dx < 396 + L)):
                return True
            if abs(xd - p) <= pw and ((dy < 116 and dy > 116 - L) or (dy > 396 and dy < 396 + L)):
                return True
        return False
    d = ImageDraw.Draw(img)
    r = 162 * f
    for y in range(int(cy - r), int(cy + r) + 1):
        for x in range(int(cx - r), int(cx + r) + 1):
            xd, yd = (x - cx) / f + 256, (y - cy) / f + 256
            if 0 <= xd <= 512 and 0 <= yd <= 512 and t(xd, yd):
                d.point((x, y), fill=_mix(STYLES["cyber"]["bg_rgb"], STYLES["cyber"]["cyan_rgb"], 235))


def _network(img):
    d = ImageDraw.Draw(img)
    S = STYLES["cyber"]
    P = [(1210, 600), (1320, 560), (1450, 600), (1560, 660), (1270, 720),
         (1400, 700), (1530, 780), (1350, 830), (1490, 860)]
    E2 = [(0, 1), (1, 2), (2, 3), (1, 4), (4, 5), (5, 2), (5, 6), (4, 7), (7, 8), (5, 8)]
    for i, j in E2:
        d.line([P[i][0], P[i][1], P[j][0], P[j][1]], fill=_mix(S["bg_rgb"], S["cyan_rgb"], 45), width=1)
    for x, y in P:
        d.ellipse([x - 7, y - 7, x + 7, y + 7], outline=_mix(S["bg_rgb"], S["cyan_rgb"], 25))
        d.rectangle([x - 2, y - 2, x + 2, y + 2], fill=_mix(S["bg_rgb"], S["cyan_rgb"], 255))
    d.line([300, 720, 1150, 720], fill=_mix(S["bg_rgb"], S["cyan_rgb"], 22), width=1)
    for x in (300, 1150):
        d.ellipse([x - 3, 717, x + 3, 723], fill=_mix(S["bg_rgb"], S["cyan_rgb"], 60))


def _traces(img):
    d = ImageDraw.Draw(img)
    S = STYLES["cyber"]
    c = _mix(S["bg_rgb"], S["cyan_rgb"], 26)
    T = [(0, 40, 420, 40), (420, 40, 420, 90), (0, 90, 300, 90), (300, 90, 300, 52),
         (W, 55, 1180, 55), (1180, 55, 1180, 100), (W, 110, 1320, 110),
         (0, 860, 460, 860), (460, 860, 460, 810), (0, 810, 330, 810), (330, 810, 330, 848),
         (W, 845, 1150, 845), (1150, 845, 1150, 800), (W, 790, 1300, 790)]
    for x0, y0, x1, y1 in T:
        d.line([x0, y0, x1, y1], fill=c, width=2)
    for x, y in [(420, 90), (300, 52), (1180, 100), (1320, 110),
                 (460, 810), (330, 848), (1150, 800), (1300, 790)]:
        d.ellipse([x - 4, y - 4, x + 4, y + 4], fill=_mix(S["bg_rgb"], S["cyan_rgb"], 60))


def _timeline(img):
    d = ImageDraw.Draw(img)
    S = STYLES["cyber"]
    d.line([250, 745, 1350, 745], fill=_mix(S["bg_rgb"], S["cyan_rgb"], 170), width=4)
    for i, x in enumerate([250, 617, 983, 1350]):
        d.ellipse([x - 9, 736, x + 9, 754], fill=_mix(S["bg_rgb"], S["cyan_rgb"], 255))
        d.ellipse([x - 16, 729, x + 16, 761], outline=_mix(S["bg_rgb"], S["cyan_rgb"], 100), width=1)
        if i == 3:
            d.ellipse([x - 26, 719, x + 26, 771], outline=_mix(S["bg_rgb"], S["cyan_rgb"], 55), width=1)
        y2 = "20%d" % (23 + i)
        f = _mono_font(24)
        # 注：与已交付版逐像素一致——宽度按默认字体计算（标签右偏约 13px 为既有观感，勿"修复"）
        w = d.textlength(y2)
        d.text((x - w / 2, 766), y2, font=f, fill=_mix(S["bg_rgb"], S["cyan_rgb"], 230))


def _accent(img):
    d = ImageDraw.Draw(img)
    S = STYLES["cyber"]
    d.line([152, 505, 560, 505], fill=_mix(S["bg_rgb"], S["cyan_rgb"], 220), width=3)
    d.rectangle([566, 501, 572, 509], fill=_mix(S["bg_rgb"], S["cyan_rgb"], 255))


def render_bg(cfg):
    """按配置渲染 1600×900 背景 PNG。
    cfg: {cover, accent, ghost, num, ghost_pos, ghost_f, num_pos, timeline}
    """
    S = STYLES["cyber"]
    img = Image.new("RGB", (W, H), S["bg_rgb"])
    d = ImageDraw.Draw(img)
    for x in range(0, W, 80):
        d.line([x, 0, x, H], fill=_mix(S["bg_rgb"], S["blue_rgb"], 10))
    for y in range(0, H, 80):
        d.line([0, y, W, y], fill=_mix(S["bg_rgb"], S["blue_rgb"], 10))
    rnd = random.Random(42)
    for _ in range(26):
        x, y = rnd.randint(0, W), rnd.randint(0, H)
        d.point((x, y), fill=_mix(S["bg_rgb"], S["cyan_rgb"], rnd.randint(20, 45)))
    _traces(img)
    if cfg.get("cover"):
        _chip(img, 1400, 170)
        _network(img)
    else:
        d.rounded_rectangle([80, 208, 1520, 838], radius=18, outline=_mix(S["bg_rgb"], S["cyan_rgb"], 45), width=2)
        d.rounded_rectangle([88, 216, 1512, 830], radius=14, outline=_mix(S["bg_rgb"], S["cyan_rgb"], 14), width=1)
        if cfg.get("ghost"):
            gp = cfg.get("ghost_pos", [1250, 440])
            gf = cfg.get("ghost_f", 1.3)
            _ghost_icon(img, cfg["ghost"], gp[0], gp[1], gf, 16)
        if cfg.get("num"):
            np_ = cfg.get("num_pos", [1150, 235])
            d.text((np_[0], np_[1]), cfg["num"], font=_mono_font(150, True), fill=_mix(S["bg_rgb"], S["cyan_rgb"], 22))
    if cfg.get("timeline"):
        _timeline(img)
    if cfg.get("accent"):
        _accent(img)
    return img


# ═══════════════════════════════════════════════════════════
# PPTX 构建（python-pptx）
# ═══════════════════════════════════════════════════════════

def _set_run(run, text, size, color, bold=False, mono=False):
    run.text = text
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = RGBColor.from_string(color)
    f.name = FONT_MONO if mono else FONT_CJK
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", FONT_CJK)


def _add_text(slide, x, y, w, h, text, size, color, bold=False, mono=False, align="left"):
    tb = slide.shapes.add_textbox(_E(x), _E(y), _E(w), _E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER}[align]
    _set_run(p.add_run(), text, size, color, bold, mono)
    return tb


def _add_box(slide, x, y, w, h, fill, line=None, lw=1.2, radius=0.10):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _E(x), _E(y), _E(w), _E(h))
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def build_pptx(output, slides, bgdir):
    """slides: [(sid, {notes, items})]；bgdir: 背景 PNG 目录（<sid>.png）"""
    S = STYLES["cyber"]
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W_EMU)
    prs.slide_height = Emu(SLIDE_H_EMU)
    blank = prs.slide_layouts[6]
    for sid, sp in slides:
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(S["bg"])
        slide.shapes.add_picture(os.path.join(bgdir, sid + ".png"), 0, 0, Emu(SLIDE_W_EMU), Emu(SLIDE_H_EMU))
        for it in sp["items"]:
            t = it["t"]
            if t in ("tag", "mono"):
                _add_text(slide, it["x"], it["y"], it["w"], it["h"], it["s"], it["size"], it["color"], False, mono=True)
            elif t == "title":
                _add_text(slide, it["x"], it["y"], it["w"], it["h"], it["s"], it["size"], it["color"], it.get("bold", True))
            elif t == "text":
                _add_text(slide, it["x"], it["y"], it["w"], it["h"], it["s"], it["size"], it["color"])
            elif t == "bullet":
                _add_text(slide, it["x"], it["y"], 44, it["h"], it["num"], 16, S["cyan"], True, mono=True)
                _add_text(slide, it["tx"], it["y"], 960, it["h"], it["text"], 18, S["body"])
            elif t == "chip":
                hl = it.get("hl", False)
                _add_box(slide, it["x"], it["y"], it["w"], it["h"],
                         S["chiphl"] if hl else S["chipf"], S["cyan"] if hl else S["chipb"], 1.5 if hl else 1.0)
                _add_text(slide, it["x"], it["y"] + 28, it["w"], 32, it["name"], 15,
                          S["cyan"] if hl else S["chipname"], True, align="center")
                _add_text(slide, it["x"], it["y"] + 68, it["w"], 26, it["field"], 11, S["dim"], align="center")
            elif t == "databar":
                for i, (num, lab) in enumerate(it["cells"]):
                    x = 110 + i * (448 + 22)
                    _add_box(slide, x, it["y"], 448, 120, S["cell"], S["cyan"], 1.2)
                    _add_text(slide, x + 26, it["y"] + 20, 400, 44, num, 24, S["cyan"], True, mono=True)
                    _add_text(slide, x + 26, it["y"] + 80, 400, 28, lab, 13, S["dim"])
            elif t == "table":
                rows = len(it["rows"])
                cols = len(it["rows"][0])
                gfx = slide.shapes.add_table(rows, cols, _E(it["x"]), _E(it["y"]), _E(it["w"]), _E(it["h"]))
                tbl = gfx.table
                tbl.first_row = False
                tbl.horz_banding = False
                sid_el = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
                if sid_el is not None:
                    sid_el.text = NO_GRID_STYLE
                for ci, cw in enumerate(it["cols"]):
                    tbl.columns[ci].width = _E(cw)
                for ri, rh in enumerate(it["rowh"]):
                    tbl.rows[ri].height = _E(rh)
                for ri, row in enumerate(it["rows"]):
                    for ci, val in enumerate(row):
                        cell = tbl.cell(ri, ci)
                        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                        cell.margin_left = _E(14)
                        cell.margin_right = _E(10)
                        cell.margin_top = _E(2)
                        cell.margin_bottom = _E(2)
                        cell.fill.solid()
                        if ri == 0:
                            cell.fill.fore_color.rgb = RGBColor.from_string(S["th"])
                            c, s, b, mo = S["cyan"], 14, True, True
                        else:
                            cell.fill.fore_color.rgb = RGBColor.from_string(S["tr0"] if ri % 2 else S["tr1"])
                            if ci == 0:
                                c, s, b, mo = S["body"], 13, False, False
                            elif ci == 1:
                                c, s, b, mo = S["cyan"], 14, True, True
                            else:
                                c, s, b, mo = S["dim"], 12, False, False
                        tf = cell.text_frame
                        tf.word_wrap = True
                        p = tf.paragraphs[0]
                        p.alignment = PP_ALIGN.LEFT
                        _set_run(p.add_run(), val, s, c, b, mo)
        ns = slide.notes_slide
        ns.notes_text_frame.text = sp.get("notes", "")
    prs.save(output)


# ═══════════════════════════════════════════════════════════
# 预览渲染（Pillow 同坐标模拟 PowerPoint）
# ═══════════════════════════════════════════════════════════

def render_previews(slides, bgdir, prevdir):
    """返回溢出警告列表。中英混排按 PowerPoint 的 latin/ea 分字渲染。"""
    S = STYLES["cyber"]

    def cjk(pt, k=0):
        try:
            return ImageFont.truetype(CJK_TTC, int(round(pt * 5 / 3)), index=[0, 2][k])
        except Exception:
            return ImageFont.load_default()

    def monof(pt, b=False):
        try:
            return ImageFont.truetype(MONO_TTC, int(round(pt * 5 / 3)), index=1 if b else 0)
        except Exception:
            return ImageFont.load_default()

    def hx(h):
        h = h.lstrip("#")
        return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))

    def tw(d, s, f):
        b = d.textbbox((0, 0), s, font=f)
        return b[2] - b[0]

    def draw_mixed(d, x, y, s, pt, color, bold=False):
        for ch in s:
            f = monof(pt, bold) if ord(ch) < 128 else cjk(pt, 1 if bold else 0)
            d.text((x, y), ch, font=f, fill=color)
            x += tw(d, ch, f)
        return x

    os.makedirs(prevdir, exist_ok=True)
    warn = []
    for sid, sp in slides:
        img = Image.open(os.path.join(bgdir, sid + ".png")).convert("RGB")
        d = ImageDraw.Draw(img)
        for it in sp["items"]:
            t = it["t"]
            if t in ("tag", "mono"):
                x = draw_mixed(d, it["x"], it["y"], it["s"], it["size"], hx(it["color"]))
                if x - it["x"] > it["w"]:
                    warn.append(f"{sid} 标签溢出: {it['s']}")
            elif t == "title":
                d.text((it["x"], it["y"]), it["s"], font=cjk(it["size"], 1), fill=hx(it["color"]))
            elif t == "text":
                d.text((it["x"], it["y"]), it["s"], font=cjk(it["size"], 0), fill=hx(it["color"]))
            elif t == "bullet":
                d.text((it["x"], it["y"]), it["num"], font=monof(16, True), fill=hx(S["cyan"]))
                f = cjk(18, 0)
                d.text((it["tx"], it["y"]), it["text"], font=f, fill=hx(S["body"]))
                if tw(d, it["text"], f) > 960:
                    warn.append(f"{sid} 要点溢出: {it['text']}")
            elif t == "chip":
                hl = it.get("hl", False)
                d.rounded_rectangle([it["x"], it["y"], it["x"] + it["w"], it["y"] + it["h"]], radius=12,
                                    fill=hx(S["chiphl"] if hl else S["chipf"]),
                                    outline=hx(S["cyan"] if hl else S["chipb"]), width=2)
                f = cjk(15, 1)
                d.text((it["x"] + (it["w"] - tw(d, it["name"], f)) / 2, it["y"] + 28), it["name"],
                       font=f, fill=hx(S["cyan"] if hl else S["chipname"]))
                f = cjk(11, 0)
                d.text((it["x"] + (it["w"] - tw(d, it["field"], f)) / 2, it["y"] + 68), it["field"],
                       font=f, fill=hx(S["dim"]))
            elif t == "databar":
                for i, (num, lab) in enumerate(it["cells"]):
                    x = 110 + i * (448 + 22)
                    d.rounded_rectangle([x, it["y"], x + 448, it["y"] + 120], radius=12,
                                        fill=hx(S["cell"]), outline=hx(S["cyan"]), width=2)
                    draw_mixed(d, x + 26, it["y"] + 20, num, 24, hx(S["cyan"]), bold=True)
                    f = cjk(13, 0)
                    d.text((x + 26, it["y"] + 80), lab, font=f, fill=hx(S["dim"]))
            elif t == "table":
                x0, y0 = it["x"], it["y"]
                cy = y0
                for ri, row in enumerate(it["rows"]):
                    rh = it["rowh"][ri]
                    cx = x0
                    for ci, val in enumerate(row):
                        cw = it["cols"][ci]
                        fill = S["th"] if ri == 0 else (S["tr0"] if ri % 2 else S["tr1"])
                        d.rectangle([cx, cy, cx + cw, cy + rh], fill=hx(fill), outline=hx(S["bg"]), width=1)
                        if ri == 0:
                            x = draw_mixed(d, cx + 14, cy + (rh - 40) // 2, val, 14, hx(S["cyan"]), bold=True)
                        elif ci == 0:
                            f = cjk(13, 0)
                            d.text((cx + 14, cy + (rh - 30) // 2), val, font=f, fill=hx(S["body"]))
                            x = cx + 14 + tw(d, val, f)
                        elif ci == 1:
                            x = draw_mixed(d, cx + 14, cy + (rh - 40) // 2, val, 14, hx(S["cyan"]), bold=True)
                        else:
                            f = cjk(12, 0)
                            d.text((cx + 14, cy + (rh - 30) // 2), val, font=f, fill=hx(S["dim"]))
                            x = cx + 14 + tw(d, val, f)
                        if x - cx > cw:
                            warn.append(f"{sid} 表格单元格溢出: {val}")
                        cx += cw
                    cy += rh
        img.save(os.path.join(prevdir, sid + ".png"))
    return warn


# ═══════════════════════════════════════════════════════════
# MCP 工具入口
# ═══════════════════════════════════════════════════════════

def _default_bg(idx, total):
    """未指定 bg 时的默认：封面(第1/末页)=cover，内容页=ghost+水印序号"""
    if idx == 1 or idx == total:
        return {"cover": True, "accent": True}
    return {"ghost": GHOSTS[(idx - 2) % len(GHOSTS)], "num": "%02d" % idx}


def write_pptx_design(output: str, spec_json: str, overwrite: bool = False) -> str:
    """按设计规格生成暗色科技风 PPTX + 预览图。

    spec_json:
    {"slides": [
        {"id": "s1", "notes": "主持稿", "bg": {"cover": true, "accent": true}, "items": [...]},
        {"id": "s2", "notes": "...", "bg": {"ghost": "target", "num": "02", "timeline": true}, "items": [...]}
    ]}
    """
    spec, err = safe_parse_json(spec_json)
    if err:
        return json.dumps({"error": f"JSON 解析失败: {err}"}, ensure_ascii=False)

    slides_spec = spec.get("slides")
    if not slides_spec or not isinstance(slides_spec, list):
        return json.dumps({"error": "必须提供 slides 数组（非空）"}, ensure_ascii=False)
    if not output:
        return json.dumps({"error": "必须指定 output 路径"}, ensure_ascii=False)
    if os.path.exists(output) and not overwrite:
        return json.dumps({"error": f"目标文件已存在: {output}（如需覆盖请传 overwrite=true）"}, ensure_ascii=False)

    # 规范化 slides：补默认 id / bg，校验 item 类型
    total = len(slides_spec)
    norm = []
    for i, sp in enumerate(slides_spec, 1):
        sid = sp.get("id") or f"s{i}"
        bg = sp.get("bg") or _default_bg(i, total)
        items = sp.get("items") or []
        if not items and not bg.get("cover"):
            return json.dumps({"error": f"第 {i} 页（{sid}）items 为空——内容页必须有条目"}, ensure_ascii=False)
        for it in items:
            if it.get("t") not in ("tag", "title", "text", "mono", "bullet", "chip", "databar", "table"):
                return json.dumps({"error": f"第 {i} 页存在未知 item 类型: {it.get('t')}"}, ensure_ascii=False)
        norm.append((sid, {"notes": sp.get("notes", ""), "items": items, "bg": bg}))

    workdir = tempfile.mkdtemp(prefix="ppt_design_")
    bgdir, prevdir = os.path.join(workdir, "bg"), os.path.join(workdir, "prev")
    os.makedirs(bgdir, exist_ok=True)
    try:
        for sid, sp in norm:
            render_bg(sp["bg"]).save(os.path.join(bgdir, sid + ".png"))
        build_pptx(output, norm, bgdir)
        warnings = render_previews(norm, bgdir, prevdir)
    except Exception as e:
        return json.dumps({"error": f"{type(e).__name__}: {e}"}, ensure_ascii=False)

    # 验证
    try:
        with zipfile.ZipFile(output) as z:
            bad = z.testzip()
            names = z.namelist()
        n_slides = len([n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")])
        n_notes = len([n for n in names if n.startswith("ppt/notesSlides/notesSlide") and n.endswith(".xml")])
        n_media = len([n for n in names if n.startswith("ppt/media/")])
    except Exception as e:
        return json.dumps({"error": f"输出文件校验失败: {e}"}, ensure_ascii=False)

    previews = sorted(os.path.join(prevdir, n) for n in os.listdir(prevdir) if n.endswith(".png"))
    result = {
        "success": True,
        "output": os.path.abspath(output),
        "slides": n_slides,
        "previews": previews,
        "warnings": warnings,
        "verification": {
            "zip_ok": bad is None,
            "notes_slides": n_notes,
            "media_files": n_media,
            "note": "相同背景图会自动去重，media_files 可少于页数；交付前必须 read_image 逐页目检 previews",
        },
    }
    if _FONT_WARN:
        result["font_warnings"] = _FONT_WARN
    return json.dumps(result, ensure_ascii=False)
